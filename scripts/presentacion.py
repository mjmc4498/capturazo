#!/usr/bin/env python3
"""Genera presentacion/capturazo.pptx con capturas reales del producto.

Requiere python-pptx y Pillow (pip install python-pptx pillow). Las imágenes salen de
presentacion/img, que llena scripts/capture-store.mjs con capturas reales de la
extensión, el celular, el visor y la videollamada de prueba.

    python3 scripts/presentacion.py
"""

from pathlib import Path

from PIL import Image, ImageDraw, ImageFilter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
IMG = ROOT / "presentacion" / "img"
OUT = ROOT / "presentacion" / "capturazo.pptx"
GEN = IMG / "generadas"
GEN.mkdir(parents=True, exist_ok=True)

BG = RGBColor(0x0B, 0x0B, 0x0C)
CARD = RGBColor(0x18, 0x18, 0x1B)
CARD2 = RGBColor(0x22, 0x22, 0x26)
TEXT = RGBColor(0xF5, 0xF5, 0xF7)
MUTED = RGBColor(0xA1, 0xA1, 0xA6)
ACCENT = RGBColor(0x1E, 0xC8, 0xFF)
FONT = "Arial"

# ---------------------------------------------------------------- imágenes

def fondo():
    path = GEN / "fondo.png"
    w, h = 1920, 1080
    base = Image.new("RGB", (w, h), (11, 11, 12))
    glow = Image.new("L", (w, h), 0)
    d = ImageDraw.Draw(glow)
    d.ellipse((1150, 520, 2250, 1400), fill=120)
    d.ellipse((-400, -500, 600, 300), fill=40)
    glow = glow.filter(ImageFilter.GaussianBlur(220))
    cyan = Image.new("RGB", (w, h), (30, 200, 255))
    base = Image.composite(cyan, base, glow.point(lambda v: int(v * 0.3)))
    base.save(path)
    return path


def redondear(src, dest, radio):
    im = Image.open(src).convert("RGBA")
    mask = Image.new("L", im.size, 0)
    ImageDraw.Draw(mask).rounded_rectangle((0, 0, im.width - 1, im.height - 1), radius=radio, fill=255)
    im.putalpha(mask)
    im.save(dest)
    return dest


def celular(nombre):
    """Pantalla del celular con bordes redondeados y marco oscuro."""
    dest = GEN / f"{nombre}-marco.png"
    pantalla = Image.open(IMG / f"{nombre}.png").convert("RGBA")
    pantalla = pantalla.resize((pantalla.width * 1170 // pantalla.width, pantalla.height * 1170 // pantalla.width), Image.LANCZOS)
    sw, sh = pantalla.size
    marco = 34
    w, h = sw + marco * 2, sh + marco * 2
    out = Image.new("RGBA", (w, h), (0, 0, 0, 0))
    cuerpo = Image.new("RGBA", (w, h), (0, 0, 0, 0))
    d = ImageDraw.Draw(cuerpo)
    d.rounded_rectangle((0, 0, w - 1, h - 1), radius=150, fill=(26, 26, 29, 255), outline=(70, 70, 76, 255), width=3)
    out.alpha_composite(cuerpo)
    mask = Image.new("L", (sw, sh), 0)
    ImageDraw.Draw(mask).rounded_rectangle((0, 0, sw - 1, sh - 1), radius=118, fill=255)
    pantalla.putalpha(mask)
    out.alpha_composite(pantalla, (marco, marco))
    out.save(dest)
    return dest


# ---------------------------------------------------------------- utilidades de diapositivas

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
FONDO = fondo()


def nueva():
    s = prs.slides.add_slide(BLANK)
    s.shapes.add_picture(str(FONDO), 0, 0, prs.slide_width, prs.slide_height)
    return s


def texto(s, x, y, w, h, contenido, size=16, color=TEXT, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=None, line=1.15):
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    parrafos = contenido if isinstance(contenido, list) else [contenido]
    for i, p_text in enumerate(parrafos):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line
        if i:
            p.space_before = Pt(size * 0.5)
        run = p.add_run()
        run.text = p_text
        f = run.font
        f.name = FONT
        f.size = Pt(size)
        f.bold = bold
        f.color.rgb = color
        if spacing is not None:
            rPr = run._r.get_or_add_rPr()
            rPr.set("spc", str(spacing))
    return box


def cabecera(s, kicker, titulo, y=0.62, w=11.8, size=36):
    texto(s, 0.7, y, 8, 0.3, kicker.upper(), size=12, color=ACCENT, bold=True, spacing=200)
    texto(s, 0.7, y + 0.38, w, 1.2, titulo, size=size, bold=True, line=1.05)


def tarjeta(s, x, y, w, h, titulo, cuerpo, numero=None, color=CARD):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = 0.08
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.shadow.inherit = False
    top = y + 0.3
    if numero is not None:
        c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.3), Inches(top), Inches(0.42), Inches(0.42))
        c.fill.solid()
        c.fill.fore_color.rgb = ACCENT
        c.line.fill.background()
        tf = c.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = str(numero)
        r.font.size = Pt(14)
        r.font.bold = True
        r.font.name = FONT
        r.font.color.rgb = RGBColor(0, 0, 0)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        top += 0.62
    # Arial negrita de 17 pt: unos 6 caracteres por pulgada. Si el título ocupa dos
    # líneas, el texto empieza más abajo.
    lineas = 2 if len(titulo) > (w - 0.6) * 6.2 else 1
    texto(s, x + 0.3, top, w - 0.6, 0.3 * lineas + 0.2, titulo, size=17, bold=True, line=1.05)
    cuerpo_y = top + 0.18 + 0.3 * lineas
    texto(s, x + 0.3, cuerpo_y, w - 0.6, h - (cuerpo_y - y) - 0.2, cuerpo, size=13, color=MUTED, line=1.25)
    return shp


def foto(s, path, x, y, w=None, h=None):
    kw = {}
    if w is not None:
        kw["width"] = Inches(w)
    if h is not None:
        kw["height"] = Inches(h)
    return s.shapes.add_picture(str(path), Inches(x), Inches(y), **kw)


def lista(s, x, y, w, items, size=15, gap=0.62):
    for i, (titulo, detalle) in enumerate(items):
        yy = y + i * gap
        dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(yy + 0.09), Inches(0.12), Inches(0.12))
        dot.fill.solid()
        dot.fill.fore_color.rgb = ACCENT
        dot.line.fill.background()
        texto(s, x + 0.3, yy, w - 0.3, 0.3, titulo, size=size, bold=True)
        texto(s, x + 0.3, yy + 0.27, w - 0.3, 0.3, detalle, size=size - 3, color=MUTED)


# ---------------------------------------------------------------- imágenes preparadas

phones = {n: celular(n) for n in ["phone-inicio", "phone-en-uso", "phone-normal", "phone-grabando", "phone-ajustes"]}
llamada = redondear(IMG / "videollamada.png", GEN / "videollamada-r.png", 36)
visor = redondear(IMG / "visor.png", GEN / "visor-r.png", 36)
popup_qr = redondear(IMG / "popup-qr.png", GEN / "popup-qr-r.png", 34)
popup_vivo = redondear(IMG / "popup-vivo.png", GEN / "popup-vivo-r.png", 34)
popup_rec = redondear(IMG / "popup-grabando.png", GEN / "popup-grabando-r.png", 34)

# ---------------------------------------------------------------- 1. portada

s = nueva()
foto(s, IMG / "logo-512.png", 0.7, 1.55, w=1.25)
texto(s, 0.7, 3.0, 7, 1.1, "Capturazo", size=66, bold=True)
texto(s, 0.7, 4.15, 6.8, 1.2, "Tu celular como cámara y micrófono de la computadora", size=24, color=MUTED, line=1.2)
texto(s, 0.7, 6.55, 7, 0.4, "Miguel Mogrovejo · Octubre de 2026", size=13, color=MUTED)
foto(s, phones["phone-en-uso"], 9.15, 0.55, h=6.4)

# ---------------------------------------------------------------- 2. el reto

s = nueva()
cabecera(s, "El reto", "Verse y oírse bien cuesta más de lo que parece")
for i, (t, c) in enumerate([
    ("Cámaras pobres", "La cámara de la laptop tiene poca resolución, mala luz y un ángulo que no favorece."),
    ("Computadoras sin cámara", "Muchas de escritorio no tienen cámara ni micrófono, y comprarlos es otro gasto."),
    ("Ruido de fondo", "El micrófono capta el teclado, la calle y el eco de la habitación."),
    ("Y el celular, guardado", "Tiene la mejor cámara y un buen micrófono, pero no hay una forma simple de usarlos."),
]):
    tarjeta(s, 0.7 + i * 3.03, 2.25, 2.83, 3.6, t, c, numero=i + 1)

# ---------------------------------------------------------------- 3. la solución

s = nueva()
cabecera(s, "La solución", "Tu celular, convertido en cámara y micrófono", w=6.6)
texto(s, 0.7, 2.55, 5.9, 1.2, "Una extensión de Chrome y una página web: escaneas un QR y tu computadora usa la cámara y el micrófono del celular como si fueran suyos.", size=16, color=MUTED, line=1.3)
lista(s, 0.7, 3.95, 5.9, [
    ("Sin instalar nada en el celular", "El QR abre la cámara en el navegador; no hay apps ni cuentas."),
    ("La cámara que quieras", "Frontal, trasera, gran angular o teleobjetivo, hasta 1080p."),
    ("Directo y cifrado", "WebRTC de punto a punto: el video no pasa por servidores."),
], gap=0.78)
foto(s, llamada, 6.75, 1.45, w=6.05)
foto(s, phones["phone-en-uso"], 6.25, 3.0, h=4.2)

# ---------------------------------------------------------------- 4. cómo se usa

s = nueva()
cabecera(s, "Cómo se usa", "Tres pasos y ya estás en cámara")
for i, (t, c) in enumerate([
    ("Escanea el código", "Haz clic en el ícono de Capturazo (o Alt+Shift+C) y apunta la cámara del celular al QR."),
    ("Toca «Empezar»", "Permite la cámara y el micrófono en el celular y elige la cámara y el modo del micrófono."),
    ("Elige Capturazo", "En tu videollamada elige «Capturazo · Cámara del celular» y «Capturazo · Micrófono del celular»."),
]):
    tarjeta(s, 0.7 + i * 2.6, 2.25, 2.45, 4.1, t, c, numero=i + 1)
foto(s, popup_qr, 8.75, 1.6, h=4.8)

# ---------------------------------------------------------------- 5. cámaras

s = nueva()
cabecera(s, "Cámaras", "Elige la cámara que mejor te muestra", w=6.2)
lista(s, 0.7, 2.35, 5.6, [
    ("Todas las cámaras del celular", "Frontal, trasera, gran angular, ultra gran angular y teleobjetivo."),
    ("Nombres claros", "También en Android, donde el sistema las llama «camera2 0, facing back»."),
    ("Girar con un toque", "Pasa de la frontal a la trasera sin cortar la llamada."),
    ("Calidad a elección", "480p, 720p o 1080p, a 24, 30 o 60 cuadros por segundo."),
    ("Linterna y espejo", "Enciende la luz del celular y mira la vista previa en espejo."),
], gap=0.8)
foto(s, phones["phone-en-uso"], 6.75, 0.6, h=6.3)
foto(s, phones["phone-ajustes"], 9.85, 0.6, h=6.3)

# ---------------------------------------------------------------- 6. micrófono

s = nueva()
cabecera(s, "Micrófono", "Normal o con reducción de ruido", w=6.2)
lista(s, 0.7, 2.35, 5.6, [
    ("Reducción de ruido", "Supresión de ruido, cancelación de eco y volumen parejo: ideal para hablar."),
    ("Normal", "Sin filtros: para música, instrumentos o el sonido del ambiente."),
    ("Cambia en vivo", "Desde el celular, el panel o el visor; la llamada sigue sin cortes."),
    ("Elige el micrófono", "El del celular o unos audífonos conectados a él."),
    ("Medidor de nivel", "Sabes en todo momento que se te está escuchando."),
], gap=0.8)
foto(s, phones["phone-en-uso"], 6.75, 0.6, h=6.3)
foto(s, phones["phone-normal"], 9.85, 0.6, h=6.3)

# ---------------------------------------------------------------- 7. videollamadas

s = nueva()
cabecera(s, "En tus videollamadas", "Aparece junto a tus cámaras", w=4.6)
lista(s, 0.7, 2.35, 4.2, [
    ("En la lista del sitio", "«Capturazo · Cámara del celular» y «Capturazo · Micrófono del celular»."),
    ("Activa sitio por sitio", "Un clic en el panel; o todos los sitios a la vez."),
    ("Sin cortes", "Si el celular se reconecta o apaga la cámara, el sitio ve una imagen de espera."),
    ("Computadoras sin cámara", "El sitio puede recibir el celular en su lugar."),
], gap=0.9)
foto(s, llamada, 5.2, 1.6, w=7.6)

# ---------------------------------------------------------------- 8. el panel

s = nueva()
cabecera(s, "El panel", "Todo bajo control desde la computadora", w=6.4)
lista(s, 0.7, 2.35, 5.8, [
    ("Vista previa en vivo", "La cámara del celular, con su resolución y cuadros por segundo."),
    ("Control remoto", "Cambia de cámara, enciende o apaga cámara y micrófono y elige el modo."),
    ("Sitios en uso", "Ves qué sitio usa el celular; el celular también lo muestra."),
    ("Varios celulares", "Elige cuál está al aire, o desconecta a cualquiera."),
], gap=0.85)
foto(s, popup_vivo, 7.35, 0.5, h=6.6)

# ---------------------------------------------------------------- 9. fotos y grabación

s = nueva()
cabecera(s, "Fotos y video", "Captura con un toque, desde donde estés", w=5.8)
lista(s, 0.7, 2.35, 5.2, [
    ("Fotos con el celular", "A la resolución de la cámara, en JPG o PNG."),
    ("Grabación con sonido", "MP4 o WebM, con el micrófono del celular."),
    ("Desde cualquier lado", "Botones en el celular, el panel y el visor."),
    ("Directo a la computadora", "Todo se guarda en Descargas › Capturazo."),
], gap=0.85)
foto(s, phones["phone-grabando"], 6.35, 0.6, h=6.3)
foto(s, popup_rec, 9.5, 1.2, h=5.0)

# ---------------------------------------------------------------- 10. visor

s = nueva()
cabecera(s, "El visor", "La cámara, en grande", w=4.6)
lista(s, 0.7, 2.35, 4.2, [
    ("Pantalla completa", "Espejo, girar y llenar la ventana."),
    ("Para OBS", "Los controles se ocultan solos: captura la ventana limpia."),
    ("Escuchar el micrófono", "Y enviarlo a otra salida de audio."),
    ("Atajos", "F pantalla completa, P foto, R grabar, M espejo, H ocultar."),
], gap=0.85)
foto(s, visor, 5.2, 1.6, w=7.6)

# ---------------------------------------------------------------- 11. cómo funciona

s = nueva()
cabecera(s, "Cómo funciona", "Una conexión directa y cifrada")


def nodo(x, y, w, h, titulo, detalle, color=CARD):
    tarjeta(s, x, y, w, h, titulo, detalle, color=color)


nodo(0.7, 2.5, 3.4, 2.2, "Celular", "Página web: abre la cámara y el micrófono elegidos y los envía por WebRTC.")
nodo(4.95, 2.5, 3.4, 2.2, "Extensión", "Recibe la señal y la reparte dentro de la computadora con conexiones locales.")
nodo(9.2, 2.5, 3.4, 2.2, "Sitios y visor", "La videollamada ve «Capturazo · Cámara del celular»; el visor la muestra en grande.")
for x in (4.15, 8.4):
    a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(3.35), Inches(0.75), Inches(0.45))
    a.fill.solid()
    a.fill.fore_color.rgb = ACCENT
    a.line.fill.background()
texto(s, 0.7, 5.0, 11.9, 0.4, "WebRTC cifrado de punto a punto · señalización PeerJS solo para el enlace inicial · retransmisión de Cloudflare si la red lo bloquea", size=13, color=MUTED, align=PP_ALIGN.CENTER)
nodo(3.2, 5.5, 6.9, 1.5, "Sin cortes por diseño", "El sitio recibe pistas propias que siguen vivas aunque el celular se reconecte o cambie de cámara.", color=CARD2)

# ---------------------------------------------------------------- 12. privacidad

s = nueva()
cabecera(s, "Privacidad", "Tu cámara, bajo tu control")
for i, (t, c) in enumerate([
    ("Sin cuentas ni analítica", "Capturazo no recopila datos personales, no muestra anuncios y no usa cookies."),
    ("Siempre a la vista", "El celular y el panel muestran qué sitio usa la cámara en cada momento."),
    ("Sitios a pedido", "El acceso es opcional: lo das sitio por sitio y lo quitas cuando quieras."),
    ("Código QR con clave", "La clave viaja en el fragmento # y se renueva en cada sesión."),
]):
    tarjeta(s, 0.7 + i * 3.03, 2.25, 2.83, 3.6, t, c, numero=i + 1)

# ---------------------------------------------------------------- 13. estado y publicación

s = nueva()
cabecera(s, "Estado · versión 1.0.0", "Listo para publicar")
for i, (t, c) in enumerate([
    ("43 comprobaciones", "Extensión y celular reales conectados, videollamada en HD con sonido, foto, grabación, visor y reconexión."),
    ("Paquete listo", "dist/capturazo-1.0.0.zip para la Chrome Web Store, con permisos mínimos y acceso a sitios opcional."),
    ("Ficha completa", "Textos, 5 capturas, 2 mosaicos, política de privacidad y justificación de cada permiso."),
]):
    tarjeta(s, 0.7 + i * 4.05, 2.25, 3.85, 2.6, t, c, numero=i + 1)
texto(s, 0.7, 5.35, 11.9, 1.2, ["Siguiente paso: publicar la página del celular en mjmc4498.github.io/capturazo y subir el ZIP 1.0.0 en el panel del desarrollador."], size=16, color=MUTED, line=1.3)

# ---------------------------------------------------------------- 14. cierre

s = nueva()
foto(s, IMG / "logo-512.png", 0.7, 1.7, w=1.25)
texto(s, 0.7, 3.15, 8, 1.1, "¡Captúralo todo!", size=56, bold=True)
texto(s, 0.7, 4.3, 7.4, 1.4, ["Escanea el QR y usa la cámara y el micrófono de tu celular en tu computadora.", "mjmc4498.github.io/capturazo"], size=20, color=MUTED, line=1.25)
texto(s, 0.7, 6.55, 7, 0.4, "Creado por Miguel Mogrovejo", size=13, color=MUTED)
foto(s, phones["phone-normal"], 9.15, 0.55, h=6.4)

prs.save(OUT)
print("  ", OUT.relative_to(ROOT), f"({len(prs.slides)} diapositivas)")
